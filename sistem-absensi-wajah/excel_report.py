"""Pembuatan rekap absensi dalam format Excel (.xlsx) via openpyxl."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
HADIR_FILL = PatternFill("solid", fgColor="E2EFDA")
TERLAMBAT_FILL = PatternFill("solid", fgColor="FFC7CE")
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def _fmt_tanggal(iso: str) -> str:
    try:
        d = datetime.strptime(iso, "%Y-%m-%d")
        hari = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"][d.weekday()]
        return f"{hari}, {d.strftime('%d/%m/%Y')}"
    except ValueError:
        return iso


def _auto_width(ws, widths: list[int]):
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt


def build_rekap(records: list[dict], users: list[dict], d_from: str, d_to: str) -> BytesIO:
    """Bangun workbook rekap dari catatan absensi (dict hasil db.list_attendance)."""
    wb = Workbook()

    # ---------- Sheet 1: detail per catatan ----------
    ws = wb.active
    ws.title = "Rekap Absensi"
    headers = ["No", "Tanggal", "Kode", "Nama Peserta", "Jam Datang", "Status Datang",
               "Jam Pulang", "Status Pulang", "Skor Kecocokan", "Bukti Foto"]

    ws.merge_cells("A1:J1")
    ws["A1"] = "REKAP ABSENSI KEHADIRAN — FACE RECOGNITION"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = CENTER
    ws.merge_cells("A2:J2")
    ws["A2"] = f"Periode: {_fmt_tanggal(d_from)} s.d. {_fmt_tanggal(d_to)}"
    ws["A2"].alignment = CENTER
    ws.merge_cells("A3:J3")
    ws["A3"] = f"Dicetak: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws["A3"].font = Font(italic=True, size=10, color="64748B")
    ws["A3"].alignment = CENTER
    ws.row_dimensions[1].height = 22

    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    row = 6
    for i, r in enumerate(records, start=1):
        values = [
            i,
            _fmt_tanggal(r["tanggal"]),
            r["kode"],
            r["nama"],
            r["waktu"],
            r["status"],
            r.get("waktu_pulang") or "-",
            r.get("status_pulang") or "-",
            round(r["skor"], 3),
            r["snapshot"] or "-",
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = BORDER
            cell.alignment = CENTER if col in (1, 2, 3, 5, 6, 7, 8) else LEFT
            if col == 9:
                cell.number_format = "0.0%"
        status_cell = ws.cell(row=row, column=6)
        status_cell.fill = HADIR_FILL if r["status"] == "Hadir" else TERLAMBAT_FILL
        pulang_cell = ws.cell(row=row, column=8)
        if r.get("status_pulang") == "Pulang Cepat":
            pulang_cell.fill = TERLAMBAT_FILL
        elif r.get("status_pulang"):
            pulang_cell.fill = HADIR_FILL
        row += 1

    ws.freeze_panes = "A6"
    _auto_width(ws, [5, 20, 10, 28, 11, 13, 11, 13, 15, 34])

    # ---------- Sheet 2: ringkasan per peserta ----------
    ws2 = wb.create_sheet("Ringkasan")
    headers2 = ["No", "Kode", "Nama Peserta", "Total Hadir", "Tepat Waktu",
                "Terlambat", "Total Absen Pulang", "Persentase Kehadiran"]
    ws2.merge_cells("A1:H1")
    ws2["A1"] = "RINGKASAN KEHADIRAN PER PESERTA"
    ws2["A1"].font = TITLE_FONT
    ws2["A1"].alignment = CENTER

    for col, title in enumerate(headers2, start=1):
        cell = ws2.cell(row=3, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    try:
        n_days = (datetime.strptime(d_to, "%Y-%m-%d") - datetime.strptime(d_from, "%Y-%m-%d")).days + 1
    except ValueError:
        n_days = max(1, len({r["tanggal"] for r in records}))

    by_user: dict[str, dict] = {}
    for r in records:
        key = r["kode"]
        agg = by_user.setdefault(key, {"kode": r["kode"], "nama": r["nama"],
                                       "hadir": 0, "tepat": 0, "terlambat": 0, "pulang": 0})
        agg["hadir"] += 1
        if r["status"] == "Terlambat":
            agg["terlambat"] += 1
        else:
            agg["tepat"] += 1
        if r.get("waktu_pulang"):
            agg["pulang"] += 1

    # peserta tanpa catatan absensi tetap tampil (kehadiran 0)
    for u in users:
        if u["kode"] not in by_user:
            by_user[u["kode"]] = {"kode": u["kode"], "nama": u["nama"],
                                  "hadir": 0, "tepat": 0, "terlambat": 0, "pulang": 0}

    row = 4
    for i, (kode, agg) in enumerate(sorted(by_user.items(), key=lambda kv: kv[1]["nama"].lower()), start=1):
        pct = agg["hadir"] / n_days if n_days > 0 else 0.0
        values = [i, kode, agg["nama"], agg["hadir"], agg["tepat"], agg["terlambat"],
                  agg["pulang"], pct]
        for col, v in enumerate(values, start=1):
            cell = ws2.cell(row=row, column=col, value=v)
            cell.border = BORDER
            cell.alignment = CENTER if col != 3 else LEFT
            if col == 8:
                cell.number_format = "0.0%"
        row += 1

    ws2.freeze_panes = "A4"
    _auto_width(ws2, [5, 10, 28, 12, 13, 12, 18, 20])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
